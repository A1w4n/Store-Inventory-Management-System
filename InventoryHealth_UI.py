"""
InventoryHealth_UI.py
──────────────────────
Inventory Health standalone sidebar page.
"""

import sys
from PySide6.QtWidgets import (
    QApplication, QWidget, QLabel, QPushButton,
    QVBoxLayout, QHBoxLayout, QFrame, QScrollArea,
    QTableWidget, QTableWidgetItem, QHeaderView,
    QAbstractItemView, QComboBox, QFileDialog, QMessageBox
)
from PySide6.QtCore import Qt, QThread, Signal
from PySide6.QtGui import QColor, QFont

from database import InventoryDatabase
from analytics_db import AnalyticsDB
from charts import HBarChart, ScatterDot, KpiCard, SectionCard



class InventoryReportWorker(QThread):
    """Generates the Inventory Health Excel report in the background."""
    finished = Signal(str)
    error    = Signal(str)

    def __init__(self, path, days, adb):
        super().__init__()
        self.path = path
        self.days = days
        self.adb  = adb

    def run(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from datetime import datetime

            wb = Workbook()

            # ── palette ───────────────────────────────────────────────────────
            PURPLE       = "4F46E5"
            PURPLE_LIGHT = "EDE9FE"
            GREEN        = "10B981"
            GREEN_LIGHT  = "D1FAE5"
            RED          = "EF4444"
            RED_LIGHT    = "FEE2E2"
            AMBER        = "F59E0B"
            AMBER_LIGHT  = "FEF3C7"
            WHITE        = "FFFFFF"
            DARK         = "111827"
            MID          = "6B7280"
            thin = Side(style="thin", color="E5E7EB")
            bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)

            def hdr_font(sz=10, color=WHITE):
                return Font(name="Arial", size=sz, bold=True, color=color)
            def body_font(sz=10, bold=False, color=DARK):
                return Font(name="Arial", size=sz, bold=bold, color=color)
            def fill(color):
                return PatternFill("solid", fgColor=color)
            def center():
                return Alignment(horizontal="center", vertical="center", wrap_text=True)
            def left():
                return Alignment(horizontal="left", vertical="center", wrap_text=True)
            def right_al():
                return Alignment(horizontal="right", vertical="center")

            def banner(ws, text, ncols):
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
                c = ws.cell(row=1, column=1, value=text)
                c.font      = Font(name="Arial", size=14, bold=True, color=WHITE)
                c.fill      = fill(PURPLE)
                c.alignment = center()
                ws.row_dimensions[1].height = 32

            def subtitle(ws, text, ncols):
                ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
                c = ws.cell(row=2, column=1, value=text)
                c.font      = Font(name="Arial", size=10, color=MID)
                c.fill      = fill(PURPLE_LIGHT)
                c.alignment = center()
                ws.row_dimensions[2].height = 16

            def table_hdr(ws, row, headers, col_start=1):
                for i, h in enumerate(headers):
                    c = ws.cell(row=row, column=col_start+i, value=h)
                    c.font = hdr_font(); c.fill = fill(PURPLE)
                    c.alignment = center(); c.border = bdr
                ws.row_dimensions[row].height = 20

            def table_row(ws, row, values, col_start=1, alt=False, row_fill=None):
                bg = row_fill or (fill("F9FAFB") if alt else fill(WHITE))
                for i, v in enumerate(values):
                    c = ws.cell(row=row, column=col_start+i, value=v)
                    c.font      = body_font()
                    c.fill      = bg
                    c.border    = bdr
                    c.alignment = right_al() if isinstance(v, (int, float)) else left()
                ws.row_dimensions[row].height = 18

            now_str = datetime.now().strftime("%B %d, %Y %H:%M")

            # ══════════════════════════════════════════════════════════════════
            # SHEET 1 — Summary KPIs
            # ══════════════════════════════════════════════════════════════════
            ws = wb.active
            ws.title = "Summary"
            ws.sheet_view.showGridLines = False
            ws.column_dimensions["A"].width = 32
            ws.column_dimensions["B"].width = 24

            banner(ws, "ProStock — Inventory Health Report", 2)
            subtitle(ws, f"Period: Last {self.days} days   |   Generated: {now_str}", 2)

            dead    = self.adb.dead_stock(self.days)
            dead_value = sum(r["tied_value"] for r in dead)
            st_rows = self.adb.sell_through_rate(self.days)
            avg_st  = 0.0
            if st_rows:
                rates  = [r["sold"] / (r["sold"] + r["stock"]) * 100 for r in st_rows]
                avg_st = sum(rates) / len(rates)

            table_hdr(ws, 4, ["Metric", "Value"])
            kpis = [
                ("Dead Stock Items",   len(dead)),
                ("Dead Stock Value",   f"₱{dead_value:,.2f}"),
                ("Avg Sell-Through",   f"{avg_st:.1f}%"),
                ("Items Tracked",      len(st_rows)),
                ("Report Period",      f"Last {self.days} days"),
            ]
            for i, (k, v) in enumerate(kpis):
                table_row(ws, 5+i, [k, v], alt=i%2==0)

            # ══════════════════════════════════════════════════════════════════
            # SHEET 2 — Sell-Through Rate
            # ══════════════════════════════════════════════════════════════════
            ws2 = wb.create_sheet("Sell-Through Rate")
            ws2.sheet_view.showGridLines = False
            for col, w in zip("ABCDE", [32, 20, 14, 14, 16]):
                ws2.column_dimensions[col].width = w

            banner(ws2, "Sell-Through Rate by Item", 5)
            subtitle(ws2, f"% of available stock sold — Last {self.days} days", 5)
            table_hdr(ws2, 3, ["Item", "Category", "Units Sold", "On Hand", "Sell-Through %"])

            for i, row in enumerate(st_rows):
                row_n = 4 + i
                rate  = row["sold"] / (row["sold"] + row["stock"]) * 100 if (row["sold"] + row["stock"]) else 0
                vals  = [row["name"], row["cat"], int(row["sold"]), int(row["stock"]), round(rate, 1)]
                table_row(ws2, row_n, vals, alt=i%2==0)
                # colour the sell-through % cell
                pct_cell = ws2.cell(row=row_n, column=5)
                if rate >= 50:
                    pct_cell.fill = fill(GREEN_LIGHT)
                    pct_cell.font = body_font(bold=True, color=GREEN)
                elif rate >= 20:
                    pct_cell.fill = fill(AMBER_LIGHT)
                    pct_cell.font = body_font(bold=True, color=AMBER)
                else:
                    pct_cell.fill = fill(RED_LIGHT)
                    pct_cell.font = body_font(bold=True, color=RED)

            # Totals
            tr = 4 + len(st_rows)
            ws2.cell(row=tr, column=1, value="TOTAL / AVG")
            ws2.cell(row=tr, column=3, value=f"=SUM(C4:C{tr-1})")
            ws2.cell(row=tr, column=4, value=f"=SUM(D4:D{tr-1})")
            ws2.cell(row=tr, column=5, value=f"=AVERAGE(E4:E{tr-1})")
            for col in range(1, 6):
                c = ws2.cell(row=tr, column=col)
                c.font = hdr_font(color=WHITE); c.fill = fill(PURPLE)
                c.border = bdr; c.alignment = right_al() if col > 1 else left()

            # ══════════════════════════════════════════════════════════════════
            # SHEET 3 — Dead Stock
            # ══════════════════════════════════════════════════════════════════
            ws3 = wb.create_sheet("Dead Stock")
            ws3.sheet_view.showGridLines = False
            for col, w in zip("ABCDE", [32, 20, 12, 16, 18]):
                ws3.column_dimensions[col].width = w

            banner(ws3, "⚠ Dead Stock — Zero Sales in Period", 5)
            subtitle(ws3, f"Items with no sales in last {self.days} days — capital tied up", 5)
            table_hdr(ws3, 3, ["Item", "Category", "Qty", "Unit Price (₱)", "Value Tied (₱)"])

            for i, row in enumerate(dead):
                row_n = 4 + i
                vals  = [row["name"], row["cat"], int(row["quantity"]),
                         float(row["price"]), float(row["tied_value"])]
                table_row(ws3, row_n, vals, alt=i%2==0,
                          row_fill=fill(RED_LIGHT) if i == 0 else None)
                # always red value column
                vc = ws3.cell(row=row_n, column=5)
                vc.font = body_font(bold=True, color=RED)
                vc.fill = fill(RED_LIGHT)

            # Totals
            tr3 = 4 + len(dead)
            ws3.cell(row=tr3, column=1, value="TOTAL")
            ws3.cell(row=tr3, column=3, value=f"=SUM(C4:C{tr3-1})")
            ws3.cell(row=tr3, column=5, value=f"=SUM(E4:E{tr3-1})")
            for col in range(1, 6):
                c = ws3.cell(row=tr3, column=col)
                c.font = hdr_font(color=WHITE); c.fill = fill(RED)
                c.border = bdr; c.alignment = right_al() if col > 1 else left()

            # ══════════════════════════════════════════════════════════════════
            # SHEET 4 — Movement Breakdown
            # ══════════════════════════════════════════════════════════════════
            ws4 = wb.create_sheet("Movement Breakdown")
            ws4.sheet_view.showGridLines = False
            for col, w in zip("AB", [28, 18]):
                ws4.column_dimensions[col].width = w

            banner(ws4, "Inventory Movement Type Breakdown", 2)
            subtitle(ws4, f"Total quantity moved per type — Last {self.days} days", 2)
            table_hdr(ws4, 3, ["Movement Type", "Total Quantity"])

            mv_rows = self.adb.movement_type_breakdown(self.days)
            mv_colors = {
                "SALE": RED, "RESTOCK": GREEN, "PURCHASE": GREEN,
                "ADJUSTMENT": AMBER, "REMOVE": RED,
            }
            for i, row in enumerate(mv_rows):
                row_n = 4 + i
                table_row(ws4, row_n, [row["movement_type"], int(row["total_qty"])], alt=i%2==0)
                color = mv_colors.get(str(row["movement_type"]).upper(), MID)
                ws4.cell(row=row_n, column=1).font = body_font(bold=True, color=color)

            tr4 = 4 + len(mv_rows)
            ws4.cell(row=tr4, column=1, value="TOTAL")
            ws4.cell(row=tr4, column=2, value=f"=SUM(B4:B{tr4-1})")
            for col in range(1, 3):
                c = ws4.cell(row=tr4, column=col)
                c.font = hdr_font(color=WHITE); c.fill = fill(PURPLE)
                c.border = bdr; c.alignment = right_al() if col > 1 else left()

            wb.save(self.path)
            self.finished.emit(self.path)

        except Exception as ex:
            self.error.emit("Report generation failed: " + str(ex))

class InventoryHealthPage(QWidget):
    """Inventory Health full sidebar page."""

    PERIOD_DAYS = {"7 Days": 7, "30 Days": 30, "90 Days": 90}

    def __init__(self, db: InventoryDatabase, parent=None):
        super().__init__(parent)
        self.adb = AnalyticsDB(db)
        self.current_days = 7
        self.setStyleSheet("background: #f9fafb;")
        self._build_ui()
        self.refresh()

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        # ── Header ──────────────────────────────────────────
        hdr_widget = QWidget()
        hdr_widget.setStyleSheet("background: #f9fafb;")
        hdr_lay = QHBoxLayout(hdr_widget)
        hdr_lay.setContentsMargins(24, 20, 24, 14)
        hdr_lay.setSpacing(10)

        title = QLabel("🏥  Inventory Health")
        title.setStyleSheet(
            "font-size: 22px; font-weight: 700; color: #111827; border: none;")
        hdr_lay.addWidget(title)
        hdr_lay.addStretch()

        self.period_combo = QComboBox()
        self.period_combo.addItems(list(self.PERIOD_DAYS.keys()))
        self.period_combo.setFixedWidth(130)
        self.period_combo.setStyleSheet("""
            QComboBox {
                background: #ffffff; border: 1.5px solid #d1d5db;
                border-radius: 8px; padding: 6px 12px;
                font-size: 12px; color: #111827; font-weight: 600;
            }
            QComboBox:hover { border-color: #4f46e5; }
            QComboBox::drop-down { border: none; padding-right: 6px; }
            QComboBox QAbstractItemView {
                background: #ffffff; border: 1px solid #d1d5db;
                color: #374151; selection-background-color: #f3f4f6;
                selection-color: #4f46e5; outline: none;
            }
        """)
        self.period_combo.currentTextChanged.connect(self._on_period_change)
        hdr_lay.addWidget(self.period_combo)

        refresh_btn = QPushButton("⟳  Refresh")
        refresh_btn.setCursor(Qt.PointingHandCursor)
        refresh_btn.setFixedHeight(36)
        refresh_btn.setStyleSheet("""
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
                    stop:0 #4f46e5, stop:1 #6366f1);
                color: white; border-radius: 8px; padding: 0 16px;
                font-size: 13px; font-weight: 600; border: none;
            }
            QPushButton:hover { background: #4338ca; }
        """)
        refresh_btn.clicked.connect(self.refresh)
        hdr_lay.addWidget(refresh_btn)

        self.report_btn = QPushButton("📊  Generate Report")
        self.report_btn.setCursor(Qt.PointingHandCursor)
        self.report_btn.setFixedHeight(36)
        self.report_btn.setStyleSheet("""
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
                    stop:0 #10b981, stop:1 #059669);
                color: white; border-radius: 8px; padding: 0 16px;
                font-size: 13px; font-weight: 600; border: none;
            }
            QPushButton:hover { background: #047857; }
            QPushButton:disabled { background: #d1d5db; color: #9ca3af; }
        """)
        self.report_btn.clicked.connect(self._generate_report)
        hdr_lay.addWidget(self.report_btn)
        root.addWidget(hdr_widget)

        line = QFrame()
        line.setFrameShape(QFrame.HLine)
        line.setStyleSheet(
            "background-color: #e5e7eb; max-height: 1px; border: none;")
        root.addWidget(line)

        # ── Scroll content ───────────────────────────────────
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        scroll.setStyleSheet("background:#f9fafb;")
        inner = QWidget()
        inner.setStyleSheet("background:#f9fafb;")
        self._lay = QVBoxLayout(inner)
        self._lay.setSpacing(16)
        self._lay.setContentsMargins(24, 20, 24, 24)
        scroll.setWidget(inner)
        root.addWidget(scroll, 1)

        # KPI row
        self.kpi_row = QHBoxLayout()
        self.kpi_row.setSpacing(12)
        self._lay.addLayout(self.kpi_row)

        # Sell-through + dead stock
        row1 = QHBoxLayout()
        row1.setSpacing(16)

        self.st_card = SectionCard("Sell-Through Rate",
                                   "% of available stock that was sold")
        self.st_table = self._make_table(
            ["Item", "Category", "Sold", "On Hand", "Sell-Through %"])
        self.st_card.add(self.st_table)
        row1.addWidget(self.st_card, 3)

        self.dead_card = SectionCard("⚠ Dead Stock",
                                     "Items with 0 sales — capital tied up")
        self.dead_table = self._make_table(
            ["Item", "Category", "Qty", "Unit Price", "Value Tied"])
        self.dead_card.add(self.dead_table)
        row1.addWidget(self.dead_card, 3)
        self._lay.addLayout(row1)

        # Scatter + movement breakdown
        row2 = QHBoxLayout()
        row2.setSpacing(16)

        self.scatter_card = SectionCard("Price vs Stock Quantity",
                                        "Each dot = one item")
        self.scatter = ScatterDot()
        self.scatter_card.add(self.scatter)
        row2.addWidget(self.scatter_card, 3)

        self.mv_card = SectionCard("Movement Type Breakdown")
        self.mv_bar = HBarChart()
        self.mv_card.add(self.mv_bar)
        row2.addWidget(self.mv_card, 2)
        self._lay.addLayout(row2)

        self._lay.addStretch()

    def _make_table(self, headers):
        t = QTableWidget(0, len(headers))
        t.setHorizontalHeaderLabels(headers)
        t.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        for i in range(1, len(headers)):
            t.horizontalHeader().setSectionResizeMode(i, QHeaderView.ResizeToContents)
        t.verticalHeader().setVisible(False)
        t.setEditTriggers(QAbstractItemView.NoEditTriggers)
        t.setSelectionBehavior(QAbstractItemView.SelectRows)
        t.setAlternatingRowColors(True)
        t.setStyleSheet("""
            QTableWidget { border:1px solid #e5e7eb; border-radius:8px;
                           gridline-color:#f3f4f6; font-size:12px; }
            QTableWidget::item { padding:6px 10px; color:#111827; }
            QTableWidget::item:selected { background:#ede9fe; color:#111827; }
            QHeaderView::section { background:#f9fafb; color:#6b7280;
                font-weight:700; font-size:11px; padding:8px 10px;
                border:none; border-bottom:1px solid #e5e7eb; }
            QTableWidget { alternate-background-color: #fafafa; }
        """)
        t.setMinimumHeight(240)
        return t

    def _on_period_change(self, text):
        self.current_days = self.PERIOD_DAYS[text]
        self.refresh()

    def _generate_report(self):
        from datetime import datetime
        default_name = f"InventoryHealth_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        path, _ = QFileDialog.getSaveFileName(
            self, "Save Inventory Health Report", default_name,
            "Excel Files (*.xlsx)"
        )
        if not path:
            return

        self.report_btn.setEnabled(False)
        self.report_btn.setText("⏳  Generating…")

        self._report_worker = InventoryReportWorker(path, self.current_days, self.adb)
        self._report_worker.finished.connect(self._on_report_done)
        self._report_worker.error.connect(self._on_report_error)
        self._report_worker.start()

    def _on_report_done(self, path):
        self.report_btn.setEnabled(True)
        self.report_btn.setText("📊  Generate Report")
        import subprocess, sys, os
        msg = QMessageBox(self)
        msg.setWindowTitle("Report Ready")
        msg.setText("✅ Report saved successfully!")
        msg.setInformativeText(path)
        msg.setStandardButtons(QMessageBox.Open | QMessageBox.Ok)
        msg.setDefaultButton(QMessageBox.Open)
        msg.setStyleSheet("QLabel { color: #111827; font-size: 13px; }")
        if msg.exec() == QMessageBox.Open:
            if sys.platform == "win32":
                os.startfile(path)
            elif sys.platform == "darwin":
                subprocess.call(["open", path])
            else:
                subprocess.call(["xdg-open", path])

    def _on_report_error(self, msg):
        self.report_btn.setEnabled(True)
        self.report_btn.setText("📊  Generate Report")
        err = QMessageBox(self)
        err.setWindowTitle("Report Error")
        err.setText("⚠️ " + msg)
        err.setStyleSheet("QLabel { color: #111827; font-size: 13px; }")
        err.exec()

    def refresh(self):
        days = self.current_days

        # KPIs
        while self.kpi_row.count():
            item = self.kpi_row.takeAt(0)
            if item.widget():
                item.widget().deleteLater()

        dead = self.adb.dead_stock(days)
        dead_value = sum(r["tied_value"] for r in dead)
        st_rows = self.adb.sell_through_rate(days)
        avg_st = 0.0
        if st_rows:
            rates = [r["sold"] / (r["sold"] + r["stock"]) * 100
                     for r in st_rows]
            avg_st = sum(rates) / len(rates)

        for icon, title, val, sub, color in [
            ("☠️", "Dead Stock Items",  str(len(dead)),
             f"₱{dead_value:,.0f} tied up",    "#ef4444"),
            ("🔄", "Avg Sell-Through",   f"{avg_st:.1f}%",
             f"Across {len(st_rows)} items",    "#10b981"),
            ("📦", "Items Tracked",      str(len(st_rows)),
             "With sales activity",             "#6366f1"),
            ("💸", "Dead Stock Value",   f"₱{dead_value:,.2f}",
             f"Over last {days} days",          "#f59e0b"),
        ]:
            self.kpi_row.addWidget(KpiCard(icon, title, val, sub, color))

        # Sell-through table
        self.st_table.setRowCount(0)
        for row in st_rows[:15]:
            r = self.st_table.rowCount()
            self.st_table.insertRow(r)
            rate = row["sold"] / (row["sold"] + row["stock"]) * 100
            color = ("#10b981" if rate >= 50 else
                     "#f59e0b" if rate >= 20 else "#ef4444")
            for col, val in enumerate([
                row["name"], row["cat"],
                str(int(row["sold"])), str(int(row["stock"])),
                f"{rate:.1f}%"
            ]):
                it = QTableWidgetItem(val)
                it.setTextAlignment(Qt.AlignVCenter |
                    (Qt.AlignLeft if col <= 1 else Qt.AlignCenter))
                if col == 4:
                    it.setForeground(QColor(color))
                    font = QFont(); font.setBold(True); it.setFont(font)
                self.st_table.setItem(r, col, it)

        # Dead stock table
        self.dead_table.setRowCount(0)
        for row in dead[:15]:
            r = self.dead_table.rowCount()
            self.dead_table.insertRow(r)
            for col, val in enumerate([
                row["name"], row["cat"],
                str(int(row["quantity"])),
                f"₱{row['price']:.2f}",
                f"₱{row['tied_value']:,.2f}",
            ]):
                it = QTableWidgetItem(val)
                it.setTextAlignment(Qt.AlignVCenter |
                    (Qt.AlignLeft if col <= 1 else Qt.AlignCenter))
                if col == 4:
                    it.setForeground(QColor("#ef4444"))
                self.dead_table.setItem(r, col, it)

        # Scatter
        CAT_COLORS = ["#6366f1", "#10b981", "#f59e0b", "#ef4444",
                      "#3b82f6", "#8b5cf6", "#ec4899", "#14b8a6"]
        pts_rows = self.adb.price_vs_quantity()
        cats = list(dict.fromkeys(r["cat"] for r in pts_rows))
        cat_color = {c: CAT_COLORS[i % len(CAT_COLORS)]
                     for i, c in enumerate(cats)}
        self.scatter.set_points([
            (float(r["price"]), float(r["quantity"]),
             r["name"], cat_color.get(r["cat"], "#6b7280"))
            for r in pts_rows
        ])

        # Movement breakdown
        mv_rows = self.adb.movement_type_breakdown(days)
        if mv_rows:
            mv_colors = {"SALE": "#ef4444", "RESTOCK": "#10b981",
                         "PURCHASE": "#10b981", "ADJUSTMENT": "#f59e0b",
                         "REMOVE": "#ef4444"}
            self.mv_bar.set_data(
                [r["movement_type"] for r in mv_rows],
                [float(r["total_qty"]) for r in mv_rows],
                [mv_colors.get(r["movement_type"].upper(), "#6b7280")
                 for r in mv_rows])


# ── Standalone runner ────────────────────────────────────────
if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    app.setFont(QFont("Segoe UI", 10))
    db = InventoryDatabase()
    win = InventoryHealthPage(db)
    win.setWindowTitle("ProStock | Inventory Health")
    win.resize(1100, 780)
    win.show()
    sys.exit(app.exec())
